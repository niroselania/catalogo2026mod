import json
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "products.json"
EMPTY_PRICES = {"withoutVat": 0, "vatPercent": 0, "withVat": 0}


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def code_part(value):
    return as_text(value).upper().replace(" ", "")


def number(value):
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return value
    return float(str(value).replace("$", "").replace(".", "").replace(",", ".").strip())


def header_name(value):
    plain = unicodedata.normalize("NFKD", as_text(value)).encode("ascii", "ignore").decode()
    return " ".join(plain.upper().split())


if len(sys.argv) != 2:
    raise SystemExit("Uso: py scripts/actualizar_precios.py RUTA_A_PRECIOS.xlsx")

price_file = Path(sys.argv[1]).resolve()
if not price_file.is_file():
    raise SystemExit(f"No se encontr\u00f3 el archivo de precios: {price_file}")

products = json.loads(CATALOG.read_text(encoding="utf-8"))
workbook = openpyxl.load_workbook(price_file, read_only=True, data_only=True)
sheet = workbook.active
headers = {header_name(value): index for index, value in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))}
required = {"SKU", "COLOR", "PRECIO SIN IVA", "IVA %", "PRECIO CON IVA"}
missing = required - headers.keys()
if missing:
    raise SystemExit(f"Faltan estas columnas: {', '.join(sorted(missing))}")

prices_by_code = {}
for row in sheet.iter_rows(min_row=2, values_only=True):
    sku = code_part(row[headers["SKU"]])
    color = code_part(row[headers["COLOR"]])
    if not sku or not color:
        continue
    without_vat = number(row[headers["PRECIO SIN IVA"]])
    vat_percent = number(row[headers["IVA %"]])
    with_vat = number(row[headers["PRECIO CON IVA"]])
    if without_vat == 0 and vat_percent == 0 and with_vat == 0:
        continue
    prices_by_code[f"{sku}{color}"] = {
        "withoutVat": without_vat,
        "vatPercent": vat_percent * 100 if 0 < vat_percent < 1 else vat_percent,
        "withVat": with_vat,
    }

catalog_codes = {product["code"] for product in products}
for product in products:
    product["prices"] = prices_by_code.get(product["code"], EMPTY_PRICES.copy())

CATALOG.write_text(json.dumps(products, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
not_in_catalog = sorted(set(prices_by_code) - catalog_codes)
print(f"Precios cargados para {len(products)} productos; {len(prices_by_code)} c\u00f3digos con precio le\u00eddos.")
print(f"Coincidentes: {len(set(prices_by_code) & catalog_codes)}. Sin coincidencia: {len(not_in_catalog)}.")
if not_in_catalog:
    print("C\u00f3digos sin coincidencia: " + ", ".join(not_in_catalog[:30]))
