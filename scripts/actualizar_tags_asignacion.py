import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "products.json"
TAG_SHEETS = ("Alto Verano", "Holidays")
TECH_TEES_SHEET = "Tech Tees"


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def code(value):
    return as_text(value).upper().replace(" ", "")


if len(sys.argv) != 2:
    raise SystemExit("Uso: py scripts/actualizar_tags_asignacion.py RUTA_A_PLANILLA.xlsx")

assignment_file = Path(sys.argv[1]).resolve()
if not assignment_file.is_file():
    raise SystemExit(f"No se encontr\u00f3 la planilla: {assignment_file}")

products = json.loads(CATALOG.read_text(encoding="utf-8"))
workbook = openpyxl.load_workbook(assignment_file, read_only=True, data_only=True)
missing_sheets = set(TAG_SHEETS + (TECH_TEES_SHEET,)) - set(workbook.sheetnames)
if missing_sheets:
    raise SystemExit(f"Faltan estas solapas: {', '.join(sorted(missing_sheets))}")

codes_by_tag = {}
for tag in TAG_SHEETS:
    codes_by_tag[tag] = {
        code(row[2])
        for row in workbook[tag].iter_rows(min_row=2, max_col=6, values_only=True)
        if code(row[2])
    }

tech_terms = {
    as_text(row[5]).lower()
    for row in workbook[TECH_TEES_SHEET].iter_rows(min_row=1, max_col=6, values_only=True)
    if as_text(row[5])
}

counts = {"Alto Verano": 0, "Tech Tees": 0, "Holidays": 0}
for product in products:
    tags = []
    if product["code"] in codes_by_tag["Alto Verano"]:
        tags.append("Alto Verano")
    if any(term in product["name"].lower() for term in tech_terms):
        tags.append("Tech Tees")
    if product["code"] in codes_by_tag["Holidays"]:
        tags.append("Holidays")
    product["tags"] = tags
    for tag in tags:
        counts[tag] += 1

CATALOG.write_text(json.dumps(products, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print("Tags actualizados: " + ", ".join(f"{tag}: {amount}" for tag, amount in counts.items()))
